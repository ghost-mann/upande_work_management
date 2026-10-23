"""Guards on sending the payroll approval pack: who gets what, and when.

This is pay data, so a wrong recipient is the worst thing that can happen. The
frappe calls are mocked; nothing here needs a site or sends mail::

    PYTHONPATH=. ~/frappe-v16-bench/env/bin/python -m unittest \\
        work_management.tests.test_payroll_pack_send -v
"""

import io
import unittest
from types import SimpleNamespace
from unittest import mock

from work_management import payroll_pack
from work_management.payroll_pack import should_send

ENTRY = "HR-PRUN-2026-00042"


class Refused(Exception):
	"""What the mocked frappe.throw raises (frappe.ValidationError's stand-in)."""


# frappe._ needs a site; the module's own `_` is swapped for the identity.
_translate = mock.patch.object(payroll_pack, "_", lambda text: text, create=True)


def setUpModule():
	_translate.start()


def tearDownModule():
	_translate.stop()

ENTRIES = [{"name": ENTRY, "start_date": "2026-10-12", "end_date": "2026-10-18",
	"workflow_state": "Pending Accounts"}]


def line(employee, name, farm, amount):
	return {"wmpay": "WMPAY-1", "employee": employee, "employee_name": name, "farm": farm,
		"task": "Picking", "block": "B1", "work_from": "2026-10-12", "work_to": "2026-10-16",
		"days": 2, "qty": 20, "rate": 10, "amount": amount, "entered_by": "a@x.com",
		"fm_approved_by": "fm@x.com", "hr_approved_by": "hr@x.com", "gm_approved_by": "gm@x.com"}


LINES = [line("E1", "Alice", "Vale", 300), line("E2", "Bob", "Endebess", 200)]
SLIPS = [
	{"entry": ENTRY, "slip": "SS1", "employee": "E1", "employee_name": "Alice", "wmpay": "WMPAY-1", "wm_amount": 300},
	{"entry": ENTRY, "slip": "SS2", "employee": "E2", "employee_name": "Bob", "wmpay": "WMPAY-1", "wm_amount": 200},
]
ROWS = (ENTRIES, LINES, SLIPS, {"E1": 2, "E2": 2}, {("E1", "Vale"): 2, ("E2", "Endebess"): 2})

SETTINGS = {
	"payroll_pack_enabled": 1,
	"payroll_pack_sender": "notifications@karenroses.com",
	"payroll_pack_to": "gm@kaitet.com\nfd@kaitet.com",
	"payroll_pack_bcc": "audit@upande.com",
}

STAGE_ROWS = [
	{"stage_label": "Actuals: Farm Manager", "scope": "Vale", "user": "fm.vale@kaitet.com"},
	{"stage_label": "Actuals: Farm Manager", "scope": "Endebess", "user": "fm.endebess@kaitet.com"},
]
ENABLED = ["fm.vale@kaitet.com", "fm.endebess@kaitet.com"]


class Sending:
	"""Patches every frappe call send_pack makes, and records the mail."""

	def __init__(self, settings=None, stage_rows=STAGE_ROWS, enabled=ENABLED, already_sent=False, rows=ROWS,
			sent_subjects=None, sender_account=True, task_pay=True):
		self.settings = dict(SETTINGS, **(settings or {}))
		self.stage_rows = stage_rows
		self.enabled = enabled
		self.already_sent = already_sent
		self.rows = rows
		# Subjects of earlier recorded sends; already_sent=True means every audience.
		self.sent_subjects = sent_subjects
		# Is there an enabled outgoing Email Account for the sender's address?
		self.sender_account = sender_account
		self.task_pay = task_pay

	def __enter__(self):
		self.patches = [
			mock.patch.object(payroll_pack, "load_rows", return_value=self.rows),
			mock.patch.object(payroll_pack, "_stage_rows", return_value=self.stage_rows),
			mock.patch.object(payroll_pack, "_enabled_users", side_effect=lambda users: [u for u in users if u in self.enabled]),
			mock.patch.object(payroll_pack, "_record", return_value="COMM-1"),
			mock.patch.object(payroll_pack, "frappe"),
			mock.patch.object(payroll_pack, "has_task_pay", side_effect=lambda name: self.task_pay),
		]
		mocks = [p.start() for p in self.patches]
		self.record, self.frappe, self.has_task_pay = mocks[3], mocks[4], mocks[5]
		self.frappe.throw.side_effect = Refused
		self.frappe.db.get_single_value.side_effect = lambda doctype, field: self.settings.get(field)
		def exists(doctype, filters):
			if doctype == "Email Account":
				return "Notifications" if self.sender_account else None
			if self.already_sent:
				return "COMM-0"
			return "COMM-0" if filters.get("subject") in (self.sent_subjects or ()) else None
		self.frappe.db.exists.side_effect = exists
		self.frappe.utils.get_url_to_form.return_value = "https://x/app/payroll-entry/" + ENTRY
		self.frappe.get_traceback.return_value = "trace"
		return self

	def __exit__(self, *exc):
		for p in self.patches:
			p.stop()

	@property
	def mails(self):
		return [c.kwargs for c in self.frappe.sendmail.call_args_list]


class TestShouldSend(unittest.TestCase):
	def test_the_one_case_that_sends(self):
		self.assertTrue(should_send("Kaitet Ltd.", True, "Pending Accounts", True, True))

	def test_each_condition_alone_stops_it(self):
		for args in [
			("Karen Roses", True, "Pending Accounts", True, True),
			("Kaitet Ltd.", False, "Pending Accounts", True, True),
			("Kaitet Ltd.", True, "Pending GM", True, True),
			("Kaitet Ltd.", True, "Pending Accounts", False, True),
			("Kaitet Ltd.", True, "Pending Accounts", True, False),
		]:
			with self.subTest(args=args):
				self.assertFalse(should_send(*args))


class FakeDoc:
	"""Like a Document: fields are attributes, and .get() is None for a missing one."""

	def __init__(self, changed=True, **fields):
		self.__dict__.update(fields)
		self._changed = changed

	def get(self, key, default=None):
		return self.__dict__.get(key, default)

	def has_value_changed(self, fieldname):
		return self._changed


class TestHook(unittest.TestCase):
	def doc(self, company="Kaitet Ltd.", state="Pending Accounts", changed=True):
		return FakeDoc(changed=changed, name=ENTRY, company=company, workflow_state=state)

	def test_a_site_without_the_workflow_saves_normally(self):
		# No Payroll Entry workflow means no workflow_state attribute at all; an
		# AttributeError here would break every after-submit save.
		doc = FakeDoc(name=ENTRY, company="Kaitet Ltd.")
		self.assertRaises(AttributeError, getattr, doc, "workflow_state")
		with mock.patch.object(payroll_pack, "frappe") as frappe, \
				mock.patch.object(payroll_pack, "has_task_pay", return_value=True) as has:
			frappe.db.get_single_value.return_value = 1
			payroll_pack.on_payroll_entry_update(doc)
		frappe.enqueue.assert_not_called()
		has.assert_not_called()

	def test_hook_enqueues_after_commit(self):
		with mock.patch.object(payroll_pack, "frappe") as frappe, \
				mock.patch.object(payroll_pack, "has_task_pay", return_value=True):
			frappe.db.get_single_value.return_value = 1
			payroll_pack.on_payroll_entry_update(self.doc())
		frappe.enqueue.assert_called_once()
		args, kwargs = frappe.enqueue.call_args
		self.assertEqual(args[0], "work_management.payroll_pack.send_pack")
		self.assertTrue(kwargs["enqueue_after_commit"])
		self.assertEqual(kwargs["entry_name"], ENTRY)
		self.assertEqual(kwargs["queue"], "long")

	def test_hook_does_nothing_for_another_state_or_when_switched_off(self):
		for doc, enabled in [(self.doc(state="Pending GM"), 1), (self.doc(), 0), (self.doc(), None),
				(self.doc(changed=False), 1), (self.doc(company="Karen Roses"), 1)]:
			with self.subTest(state=doc.workflow_state, enabled=enabled), \
					mock.patch.object(payroll_pack, "frappe") as frappe, \
					mock.patch.object(payroll_pack, "has_task_pay", return_value=True):
				frappe.db.get_single_value.return_value = enabled
				payroll_pack.on_payroll_entry_update(doc)
			frappe.enqueue.assert_not_called()

	def test_hook_does_not_query_slips_unless_the_rest_holds(self):
		with mock.patch.object(payroll_pack, "frappe") as frappe, \
				mock.patch.object(payroll_pack, "has_task_pay", return_value=True) as has, \
				mock.patch.object(payroll_pack, "should_send", wraps=payroll_pack.should_send) as decide:
			frappe.db.get_single_value.return_value = 1
			payroll_pack.on_payroll_entry_update(self.doc(state="Pending GM"))
		has.assert_not_called()
		decide.assert_called_once()
		with mock.patch.object(payroll_pack, "frappe") as frappe, \
				mock.patch.object(payroll_pack, "has_task_pay", return_value=True) as has, \
				mock.patch.object(payroll_pack, "should_send", wraps=payroll_pack.should_send) as decide:
			frappe.db.get_single_value.return_value = 1
			payroll_pack.on_payroll_entry_update(self.doc())
		has.assert_called_once_with(ENTRY)
		decide.assert_called_once()
		frappe.enqueue.assert_called_once()

	def test_hook_does_nothing_without_task_pay(self):
		with mock.patch.object(payroll_pack, "frappe") as frappe, \
				mock.patch.object(payroll_pack, "has_task_pay", return_value=False):
			frappe.db.get_single_value.return_value = 1
			payroll_pack.on_payroll_entry_update(self.doc())
		frappe.enqueue.assert_not_called()


class TestSendPack(unittest.TestCase):
	def test_first_send_is_not_revised(self):
		with Sending() as s:
			payroll_pack.send_pack(ENTRY)
		full = s.mails[0]
		self.assertEqual(full["subject"], f"Task pay for approval — {ENTRY} — KES 500")
		self.assertEqual(full["recipients"], ["gm@kaitet.com", "fd@kaitet.com"])
		self.assertEqual(full["bcc"], ["audit@upande.com"])
		self.assertEqual(full["sender"], "notifications@karenroses.com")
		self.assertEqual(full["reference_doctype"], "Payroll Entry")
		self.assertEqual(full["reference_name"], ENTRY)
		self.assertEqual(full["attachments"][0]["fname"], "Task-pay-2026-10-12-to-2026-10-18.xlsx")
		self.assertIsInstance(full["attachments"][0]["fcontent"], bytes)

	def test_second_send_is_revised(self):
		with Sending(already_sent=True) as s:
			payroll_pack.send_pack(ENTRY)
		self.assertTrue(all(m["subject"].startswith("Revised: Task pay for approval") for m in s.mails))
		# The check is the brief's: same entry, the pack's subject, and not a test.
		# Exact stored subjects, no LIKE: names come from the database.
		comms = [c for c in s.frappe.db.exists.call_args_list if c.args[0] == "Communication"]
		subjects = [c.args[1]["subject"] for c in comms]
		self.assertEqual(subjects, [f"Task pay for approval \u2014 {ENTRY}",
			f"Task pay for approval \u2014 {ENTRY} \u2014 Endebess",
			f"Task pay for approval \u2014 {ENTRY} \u2014 Vale"])
		for c in comms:
			self.assertEqual(c.args[1]["reference_doctype"], "Payroll Entry")
			self.assertEqual(c.args[1]["reference_name"], ENTRY)
			self.assertEqual(c.args[1]["communication_type"], "Communication")
			self.assertNotIn("like", str(c.args[1]))

	def test_each_farm_manager_gets_only_their_farm(self):
		with Sending() as s:
			result = payroll_pack.send_pack(ENTRY)
		by_to = {tuple(m["recipients"]): m for m in s.mails}
		vale = by_to[("fm.vale@kaitet.com",)]
		endebess = by_to[("fm.endebess@kaitet.com",)]
		self.assertIn("Vale", vale["subject"])
		self.assertIn("KES 300", vale["subject"])
		self.assertNotIn("Bob", vale["message"])
		self.assertIn("Endebess", endebess["subject"])
		self.assertIn("KES 200", endebess["subject"])
		self.assertFalse(vale.get("bcc"))
		self.assertFalse(endebess.get("bcc"))
		self.assertEqual(vale["attachments"][0]["fname"], "Task-pay-2026-10-12-to-2026-10-18-Vale.xlsx")
		self.assertEqual(len(s.mails), 3)
		self.assertEqual(len(result["sent"]), 3)

	def test_revised_is_decided_per_audience(self):
		# Only Vale's pack went before: Vale's is revised, the full pack and Endebess are not.
		with Sending(sent_subjects={f"Task pay for approval \u2014 {ENTRY} \u2014 Vale"}) as s:
			payroll_pack.send_pack(ENTRY)
		revised = {m["recipients"][0]: m["subject"].startswith("Revised: ") for m in s.mails}
		self.assertEqual(revised, {"gm@kaitet.com": False, "fm.vale@kaitet.com": True,
			"fm.endebess@kaitet.com": False})
		# A farm record never makes the full pack revised, and vice versa.
		with Sending(sent_subjects={f"Task pay for approval \u2014 {ENTRY}"}) as s:
			payroll_pack.send_pack(ENTRY)
		revised = {m["recipients"][0]: m["subject"].startswith("Revised: ") for m in s.mails}
		self.assertEqual(revised, {"gm@kaitet.com": True, "fm.vale@kaitet.com": False,
			"fm.endebess@kaitet.com": False})

	def test_real_sends_are_recorded_per_audience(self):
		with Sending() as s:
			payroll_pack.send_pack(ENTRY)
		self.assertEqual([c.args for c in s.record.call_args_list], [
			(ENTRY, None, ["gm@kaitet.com", "fd@kaitet.com"]),
			(ENTRY, "Endebess", ["fm.endebess@kaitet.com"]),
			(ENTRY, "Vale", ["fm.vale@kaitet.com"]),
		])
		for m in s.mails:
			# Not linked: a linked Communication would put the pay in the timeline.
			self.assertNotIn("communication", m)
			# An approver who unsubscribed from some other mail must still get this.
			self.assertEqual(m["add_unsubscribe_link"], 0)

	def test_blank_sender_sends_nothing_and_logs(self):
		with Sending(settings={"payroll_pack_sender": ""}) as s:
			result = payroll_pack.send_pack(ENTRY)
		s.frappe.sendmail.assert_not_called()
		s.record.assert_not_called()
		s.frappe.log_error.assert_called()
		self.assertTrue(result["skipped"])

	def test_a_sender_with_no_outgoing_account_sends_nothing_and_logs(self):
		# frappe would fall back to the default outgoing account -- never for pay data.
		with Sending(sender_account=False) as s:
			result = payroll_pack.send_pack(ENTRY)
		s.frappe.sendmail.assert_not_called()
		s.record.assert_not_called()
		s.frappe.db.exists.assert_any_call("Email Account",
			{"email_id": "notifications@karenroses.com", "enable_outgoing": 1})
		self.assertEqual(s.frappe.log_error.call_args.kwargs["title"],
			"Payroll approval pack sender has no outgoing account")
		self.assertTrue(any("outgoing" in x for x in result["skipped"]))

	def test_a_sender_with_no_outgoing_account_refuses_a_test_send(self):
		with Sending(sender_account=False) as s:
			with self.assertRaises(Refused):
				payroll_pack.send_pack(ENTRY, ["austin@upande.com"])
		s.frappe.sendmail.assert_not_called()
		self.assertIn("notifications@karenroses.com", s.frappe.throw.call_args.args[0])
		self.assertIn("outgoing Email Account", s.frappe.throw.call_args.args[0])

	def test_switched_off_after_queueing_sends_nothing(self):
		with Sending(settings={"payroll_pack_enabled": 0}) as s:
			result = payroll_pack.send_pack(ENTRY)
		s.frappe.sendmail.assert_not_called()
		s.record.assert_not_called()
		s.frappe.log_error.assert_called_once()
		self.assertEqual(s.frappe.log_error.call_args.kwargs["title"], "Payroll approval pack not sent")
		self.assertTrue(result["skipped"])

	def test_task_pay_gone_after_queueing_sends_nothing(self):
		# The slips were redone without task pay: a "KES 0" pack must not go out.
		with Sending(task_pay=False) as s:
			result = payroll_pack.send_pack(ENTRY)
		s.has_task_pay.assert_called_once_with(ENTRY)
		s.frappe.sendmail.assert_not_called()
		s.record.assert_not_called()
		s.frappe.log_error.assert_called_once()
		self.assertTrue(result["skipped"])

	def test_a_test_send_does_not_need_the_switch_or_task_pay_rechecked(self):
		with Sending(settings={"payroll_pack_enabled": 0}, task_pay=False) as s:
			payroll_pack.send_pack(ENTRY, ["austin@upande.com"])
		self.assertEqual(len(s.mails), 3)
		s.has_task_pay.assert_not_called()

	def test_blank_sender_on_a_test_send_raises(self):
		with Sending(settings={"payroll_pack_sender": None}) as s:
			s.frappe.throw.side_effect = ValueError("no sender")
			with self.assertRaises(ValueError):
				payroll_pack.send_pack(ENTRY, ["austin@upande.com"])
		s.frappe.sendmail.assert_not_called()

	def test_test_send_goes_only_to_test_recipients(self):
		test = ["hro@endebesscoffee.com", "austin@upande.com"]
		with Sending() as s:
			result = payroll_pack.send_pack(ENTRY, test)
		self.assertEqual(len(s.mails), 3)  # full pack + Vale + Endebess
		for m in s.mails:
			self.assertEqual(m["recipients"], test)
			self.assertFalse(m.get("bcc"))
			self.assertTrue(m["subject"].startswith("TEST: "))
			self.assertNotIn("communication", m)
		s.record.assert_not_called()
		self.assertTrue(all(r["to"] == test for r in result["sent"]))

	def test_an_empty_test_list_sends_nothing_rather_than_the_real_list(self):
		with Sending() as s:
			payroll_pack.send_pack(ENTRY, [])
		s.frappe.sendmail.assert_not_called()

	def test_farm_without_fm_is_named_in_the_full_body_and_gets_no_mail(self):
		with Sending(stage_rows=STAGE_ROWS[:1]) as s:
			result = payroll_pack.send_pack(ENTRY)
		self.assertEqual(len(s.mails), 2)
		self.assertIn("No farm manager is set up for: Endebess", s.mails[0]["message"])
		self.assertNotIn("fm.endebess@kaitet.com", str(s.mails))
		self.assertTrue(any("Endebess" in str(x) for x in result["skipped"]))

	def test_a_disabled_farm_manager_gets_no_mail(self):
		with Sending(enabled=["fm.vale@kaitet.com"]) as s:
			payroll_pack.send_pack(ENTRY)
		self.assertNotIn("fm.endebess@kaitet.com", str(s.mails))
		self.assertIn("No farm manager is set up for: Endebess", s.mails[0]["message"])

	def test_no_approver_list_sends_no_full_pack_and_logs(self):
		with Sending(settings={"payroll_pack_to": ""}) as s:
			payroll_pack.send_pack(ENTRY)
		self.assertNotIn(["gm@kaitet.com", "fd@kaitet.com"], [m["recipients"] for m in s.mails])
		self.assertFalse(any(m.get("bcc") for m in s.mails))
		self.assertTrue(all(m["recipients"][0].startswith("fm.") for m in s.mails))
		s.frappe.log_error.assert_called()

	def test_a_failure_is_logged_and_swallowed_on_a_real_send(self):
		with Sending() as s:
			s.frappe.sendmail.side_effect = RuntimeError("smtp down")
			payroll_pack.send_pack(ENTRY)
		s.frappe.log_error.assert_called_with(title="Payroll approval pack failed", message="trace")

	def test_a_failure_is_raised_on_a_test_send(self):
		with Sending() as s:
			s.frappe.sendmail.side_effect = RuntimeError("smtp down")
			with self.assertRaises(RuntimeError):
				payroll_pack.send_pack(ENTRY, ["austin@upande.com"])
		s.frappe.log_error.assert_called_with(title="Payroll approval pack failed", message="trace")


class TestRecord(unittest.TestCase):
	"""The Communication left on the Payroll Entry for each real send."""

	def record(self, farm):
		with mock.patch.object(payroll_pack, "frappe") as frappe:
			frappe.get_doc.return_value.name = "COMM-9"
			name = payroll_pack._record(ENTRY, farm, ["gm@kaitet.com", "fd@kaitet.com"])
		self.assertEqual(name, "COMM-9")
		frappe.get_doc.return_value.insert.assert_called_once_with(ignore_permissions=True)
		return frappe.get_doc.call_args.args[0]

	def test_record_is_not_email_so_it_gets_no_timeline_links(self):
		# communication_medium "Email" makes Communication.validate link every
		# address to its Contact (and on to Leads...), whose timelines would then
		# show this record to anyone who can read that Contact.
		for farm in (None, "Vale"):
			doc = self.record(farm)
			self.assertNotEqual(doc["communication_medium"], "Email")
			self.assertEqual(doc["communication_medium"], "Other")
			self.assertNotIn("timeline_links", doc)
			self.assertNotIn("bcc", doc)
			self.assertNotIn("cc", doc)
			self.assertEqual((doc["reference_doctype"], doc["reference_name"]), ("Payroll Entry", ENTRY))

	def test_record_holds_no_figures(self):
		full, farm = self.record(None), self.record("Vale")
		self.assertEqual(full["subject"], f"Task pay for approval \u2014 {ENTRY}")
		self.assertEqual(farm["subject"], f"Task pay for approval \u2014 {ENTRY} \u2014 Vale")
		for doc in (full, farm):
			text = str(doc)
			self.assertNotIn("KES", text)
			self.assertNotRegex(doc["content"], r"\d")
			self.assertNotIn("<", doc["content"])


class TestSendTestPack(unittest.TestCase):
	def test_only_system_managers_and_addresses_are_split(self):
		with mock.patch.object(payroll_pack, "frappe") as frappe, \
				mock.patch.object(payroll_pack, "send_pack", return_value={"sent": [], "skipped": []}) as send:
			payroll_pack.send_test_pack(ENTRY, "a@x.com, b@x.com,a@x.com")
		frappe.only_for.assert_called_once_with("System Manager")
		send.assert_called_once_with(ENTRY, ["a@x.com", "b@x.com"])

	def test_no_addresses_is_refused(self):
		with mock.patch.object(payroll_pack, "frappe") as frappe, \
				mock.patch.object(payroll_pack, "send_pack") as send:
			frappe.throw.side_effect = ValueError("no recipients")
			with self.assertRaises(ValueError):
				payroll_pack.send_test_pack(ENTRY, " , ")
		send.assert_not_called()

	def run_failing(self, error):
		with mock.patch.object(payroll_pack, "frappe") as frappe, \
				mock.patch.object(payroll_pack, "send_pack", side_effect=error):
			frappe.ValidationError, frappe.PermissionError = Refused, PermissionError
			frappe.throw.side_effect = Refused
			try:
				payroll_pack.send_test_pack(ENTRY, "a@x.com")
			finally:
				self.frappe = frappe

	def test_an_unexpected_error_becomes_a_generic_message(self):
		# Uncaught, frappe's request handler would snapshot the local variables
		# (pay figures) into an Error Log. send_pack has already logged the traceback.
		with self.assertRaises(Refused):
			self.run_failing(KeyError("slip_total"))
		self.frappe.throw.assert_called_once_with(
			"Could not build the pay breakdown. The error has been logged.")

	def test_validation_and_permission_errors_pass_through(self):
		for error in (Refused("Set the sender"), PermissionError("no")):
			with self.subTest(error=error), self.assertRaises(type(error)) as caught:
				self.run_failing(error)
			self.assertIs(caught.exception, error)
			self.frappe.throw.assert_not_called()


class TestLoadRows(unittest.TestCase):
	def test_queries_are_parameterised_and_shapes_match_the_builder(self):
		calls = []

		def sql(query, values=None, as_dict=False):
			calls.append((query, values))
			q = " ".join(query.split())
			if "from `tabPayroll Entry`" in q:
				return [dict(ENTRIES[0])]
			if "from `tabSalary Slip`" in q:
				return [dict(SLIPS[0]), {"entry": ENTRY, "slip": "SS3", "employee": "E9",
					"employee_name": "Zed", "wmpay": None, "wm_amount": None}]
			if "from `tabWork Payment Line`" in q:
				return [dict(LINES[0])]
			if "group by wae.employee, wma.farm" in q:
				return [{"employee": "E1", "farm": "Vale", "days": 2}]
			if "from `tabWork Actuals Employee`" in q:
				return [{"employee": "E1", "days": 2}]
			raise AssertionError(q)

		with mock.patch.object(payroll_pack, "frappe") as frappe:
			frappe.db.sql.side_effect = sql
			entries, lines, slips, days, days_farm = payroll_pack.load_rows([ENTRY])
		self.assertEqual(entries[0]["name"], ENTRY)
		self.assertEqual(slips[1]["wm_amount"], 0)
		self.assertEqual(days, {"E1": 2})
		self.assertEqual(days_farm, {("E1", "Vale"): 2})
		line_query = " ".join(next(q for q, _ in calls if "`tabWork Payment Line`" in q).split())
		self.assertIn("left join `tabTask` t on t.name = wpl.task", line_query)
		self.assertIn("wpl.task as task_id", line_query)
		for query, values in calls:
			self.assertNotIn(ENTRY, query)
			self.assertNotIn("WMPAY-1", query)
			self.assertIsNotNone(values)

	def test_no_slips_means_no_line_queries(self):
		with mock.patch.object(payroll_pack, "frappe") as frappe:
			frappe.db.sql.side_effect = lambda q, v=None, as_dict=False: (
				[dict(ENTRIES[0])] if "`tabPayroll Entry`" in q else [])
			rows = payroll_pack.load_rows([ENTRY])
		self.assertEqual(rows[1:], ([], [], {}, {}))
		self.assertEqual(frappe.db.sql.call_count, 2)


class TestRangeDownload(unittest.TestCase):
	"""The breakdown for a date range: a second way out for pay data."""

	def download(self, from_date, to_date, names=(ENTRY,), denied=(), build_error=None):
		response = SimpleNamespace()
		with mock.patch.object(payroll_pack, "frappe") as frappe, \
				mock.patch.object(payroll_pack, "entries_in_range", return_value=list(names)) as found, \
				mock.patch.object(payroll_pack, "load_rows", return_value=ROWS) as load:
			frappe.response = response
			frappe.throw.side_effect = Refused
			frappe.PermissionError = PermissionError
			frappe.ValidationError = Refused
			frappe.get_traceback.return_value = "trace"
			if build_error:
				load.side_effect = build_error

			def has_permission(doctype, ptype, doc=None, throw=False):
				if doc in denied or (doc is None and None in denied):
					raise PermissionError(doc)
				return True
			frappe.has_permission.side_effect = has_permission
			try:
				payroll_pack.download_breakdown(from_date, to_date)
			finally:
				self.frappe, self.found, self.load, self.response = frappe, found, load, response
		return response

	def test_range_longer_than_92_days_is_refused(self):
		# 92 days counted inclusively is the most; 93 is refused before any query.
		self.download("2026-07-24", "2026-10-23")
		with self.assertRaises(Refused):
			self.download("2026-07-23", "2026-10-23")
		self.found.assert_not_called()
		self.assertFalse(hasattr(self.response, "filecontent"))

	def test_from_after_to_is_refused(self):
		with self.assertRaises(Refused):
			self.download("2026-10-19", "2026-10-18")
		self.found.assert_not_called()
		self.assertFalse(hasattr(self.response, "filecontent"))

	def test_a_missing_date_is_refused_not_read_as_today(self):
		for args in (("", "2026-10-18"), ("2026-10-12", None)):
			with self.subTest(args=args), self.assertRaises(Refused):
				self.download(*args)
			self.found.assert_not_called()

	def test_empty_range_says_no_payrolls(self):
		with self.assertRaises(Refused):
			self.download("2026-10-01", "2026-10-07", names=())
		self.frappe.throw.assert_called_once_with("No task-pay payrolls in that range.")
		self.load.assert_not_called()
		self.assertFalse(hasattr(self.response, "filecontent"))

	def test_range_with_an_unreadable_entry_is_refused(self):
		with self.assertRaises(PermissionError):
			self.download("2026-10-01", "2026-10-31", names=(ENTRY, "HR-PRUN-2026-00043"),
				denied=("HR-PRUN-2026-00043",))
		self.load.assert_not_called()
		self.assertFalse(hasattr(self.response, "filecontent"))
		self.assertFalse(hasattr(self.response, "filename"))

	def test_every_entry_is_checked_with_throw(self):
		names = (ENTRY, "HR-PRUN-2026-00043")
		self.download("2026-10-01", "2026-10-31", names=names)
		self.assertEqual(self.frappe.has_permission.call_args_list,
			[mock.call("Payroll Entry", "read", throw=True)]
			+ [mock.call("Payroll Entry", "read", doc=n, throw=True) for n in names])
		self.load.assert_called_once_with(list(names))

	def test_no_payroll_access_is_refused_before_any_query(self):
		# Otherwise "No task-pay payrolls in that range" tells anyone whether payrolls exist.
		with self.assertRaises(PermissionError):
			self.download("2026-10-01", "2026-10-31", denied=(None,))
		self.assertEqual(self.frappe.has_permission.call_args_list[0],
			mock.call("Payroll Entry", "read", throw=True))
		self.found.assert_not_called()
		self.frappe.db.sql.assert_not_called()
		self.frappe.throw.assert_not_called()

	def test_an_unexpected_error_is_logged_without_locals_and_reported_generically(self):
		with self.assertRaises(Refused):
			self.download("2026-10-12", "2026-10-18", build_error=KeyError("slip_total"))
		self.frappe.log_error.assert_called_once_with(
			title="Payroll pay breakdown download failed", message="trace")
		self.frappe.throw.assert_called_once_with(
			"Could not build the pay breakdown. The error has been logged.")
		self.assertFalse(hasattr(self.response, "filecontent"))

	def test_validation_errors_in_the_build_pass_through(self):
		error = Refused("from inside")
		with self.assertRaises(Refused) as caught:
			self.download("2026-10-12", "2026-10-18", build_error=error)
		self.assertIs(caught.exception, error)
		self.frappe.log_error.assert_not_called()

	def test_range_download_sets_a_binary_xlsx_response(self):
		response = self.download("2026-10-12", "2026-10-18")
		self.found.assert_called_once()
		self.assertEqual(response.type, "binary")
		self.assertEqual(response.filename, "Task-pay-2026-10-12-to-2026-10-18.xlsx")
		self.assertEqual(response.filecontent[:2], b"PK")
		from openpyxl import load_workbook
		wb = load_workbook(io.BytesIO(response.filecontent))
		self.assertEqual(wb.sheetnames[0], "Payrolls")

	def test_download_allows_get(self):
		# The button opens a URL, so this one method must accept GET.
		import frappe
		self.assertIn("GET", frappe.allowed_http_methods_for_whitelisted_func[payroll_pack.download_breakdown])
		self.assertNotIn(payroll_pack.download_breakdown, frappe.guest_methods)


class TestEntriesInRange(unittest.TestCase):
	def run_query(self, has_state=True, task_pay=("A",)):
		with mock.patch.object(payroll_pack, "frappe") as frappe, \
				mock.patch.object(payroll_pack, "has_task_pay", side_effect=lambda n: n in task_pay):
			frappe.db.has_column.return_value = has_state
			frappe.db.sql.return_value = [("A",), ("B",)]
			names = payroll_pack.entries_in_range("2026-10-01", "2026-10-31")
		query, values = frappe.db.sql.call_args.args
		return names, " ".join(query.split()), values

	def test_overlap_company_docstatus_and_state_are_parameterised(self):
		names, query, values = self.run_query()
		self.assertEqual(names, ["A"])
		self.assertIn("start_date <= %(to_date)s", query)
		self.assertIn("end_date >= %(from_date)s", query)
		self.assertIn("company = %(company)s", query)
		self.assertIn("docstatus = 1", query)
		# A NULL state (an entry from before the workflow) is not a dead one.
		self.assertIn("(workflow_state is null or workflow_state not in %(dead)s)", query)
		self.assertIn("order by start_date", query)
		self.assertNotIn("2026-10", query)
		self.assertNotIn("Kaitet", query)
		self.assertEqual(values["company"], "Kaitet Ltd.")
		self.assertEqual(set(values["dead"]), {"Cancelled", "Rejected"})

	def test_without_a_workflow_column_the_state_filter_is_left_out(self):
		_, query, _ = self.run_query(has_state=False)
		self.assertNotIn("workflow_state", query)


if __name__ == "__main__":
	unittest.main()
