"""Guards on the payroll approval pack: what an approver sees before pay is released.

Pure: rows in, a pack out. No site::

    ./env/bin/python -m unittest work_management.tests.test_payroll_pack_book -v
"""

import copy
import datetime
import io
import unittest

import openpyxl

from work_management.payroll_pack_book import (
	attachment_name, build_pack, email_html, farm_recipients, farms_in, filter_farm,
	split_addresses, subject, workbook_bytes,
)


def line(wmpay, employee, name, farm, task, days, qty, rate, amount):
	return {"wmpay": wmpay, "employee": employee, "employee_name": name, "farm": farm,
		"task": task, "block": "B1", "work_from": "2026-10-12", "work_to": "2026-10-16",
		"days": days, "qty": qty, "rate": rate, "amount": amount, "entered_by": "a@x.com",
		"fm_approved_by": "fm@x.com", "hr_approved_by": "hr@x.com", "gm_approved_by": "gm@x.com"}


def slip(entry, slip_name, employee, name, wmpay, wm_amount):
	return {"entry": entry, "slip": slip_name, "employee": employee, "employee_name": name,
		"wmpay": wmpay, "wm_amount": wm_amount}


ENTRIES = [
	{"name": "PE1", "start_date": "2026-10-12", "end_date": "2026-10-18", "workflow_state": "Accounts"},
	{"name": "PE2", "start_date": "2026-10-19", "end_date": "2026-10-25", "workflow_state": "Accounts"},
]

# E1 works two farms; E2 is paid 500 for 450 of work; E3 is on two slips in PE1 (and a
# third in PE2, which is just next week's pay); E9 has no work.
LINES = [
	line("P1", "E1", "Alice", "Endebess", "Picking", 3, 30, 10, 300),
	line("P2", "E1", "Alice", "Vale", "Weeding", 2, 20, 10, 200),
	line("P3", "E2", "Bob", "Endebess", "Picking", 4, 45, 10, 450),
	line("P4", "E3", "Carol", "Vale", "Weeding", 2, 40, 10, 400),
	line("P5", "E3", "Carol", "Vale", "Pruning", 3, 55, 10, 550),
]
SLIPS = [
	slip("PE1", "S1", "E1", "Alice", "P1", 300),
	slip("PE1", "S1", "E1", "Alice", "P2", 200),
	slip("PE1", "S2", "E2", "Bob", "P3", 500),
	slip("PE1", "S3", "E3", "Carol", "P4", 400),
	slip("PE1", "S5", "E3", "Carol", None, 0),
	slip("PE2", "S4", "E3", "Carol", "P5", 550),
	slip("PE1", "S9", "E9", "Zed", None, 0),
]
DAYS = {"E1": 4, "E2": 4, "E3": 5, "E9": 1}


def pack(lines=None, slips=None, days=None, days_farm=None):
	return build_pack(copy.deepcopy(ENTRIES), copy.deepcopy(LINES if lines is None else lines),
		copy.deepcopy(SLIPS if slips is None else slips), DAYS if days is None else days, days_farm)


def by_emp(rows):
	return {r["employee"]: r for r in rows}


def kinds(p):
	return {(c["employee"], c["kind"]) for c in p["checks"]}


class TestHeadlineFigures(unittest.TestCase):

	def test_days_worked_counts_calendar_days_not_line_days(self):
		"""E1's lines add to 5 days but they worked 4 calendar days."""
		self.assertEqual(by_emp(pack()["workers"])["E1"]["days_worked"], 4)

	def test_headline_total_is_the_slip_total(self):
		t = pack()["totals"]
		self.assertEqual(t["slip_total"], 1950.0)
		self.assertEqual(t["line_total"], 1900.0)

	def test_man_days_is_the_sum_of_days_worked(self):
		self.assertEqual(pack()["totals"]["man_days"], 14)

	def test_workers_sort_by_first_farm_then_name(self):
		order = [w["employee"] for w in pack()["workers"]]
		self.assertEqual(order, ["E1", "E2", "E3", "E9"])

	def test_entries_carry_their_worker_count_and_amount(self):
		e = {r["name"]: r for r in pack()["entries"]}
		self.assertEqual((e["PE1"]["workers"], e["PE1"]["amount"]), (4, 1400.0))
		self.assertEqual((e["PE2"]["workers"], e["PE2"]["amount"]), (1, 550.0))


class TestChecks(unittest.TestCase):

	def test_a_slip_with_no_work_is_a_check(self):
		p = pack()
		self.assertIn(("E9", "no_work"), kinds(p))
		w = by_emp(p["workers"])["E9"]
		self.assertEqual((w["farms"], w["days_worked"], w["tasks"], w["wm_amount"]), ([], 1, 0, 0))

	def test_slip_and_lines_disagreeing_is_a_check(self):
		p = pack()
		w = by_emp(p["workers"])["E2"]
		self.assertEqual((w["slip_amount"], w["wm_amount"], w["difference"]), (500.0, 450.0, 50.0))
		self.assertIn(("E2", "amount_mismatch"), kinds(p))
		self.assertEqual(w["flag"], "amount_mismatch")

	def test_a_worker_on_two_slips_in_one_entry_is_a_check(self):
		p = pack()
		self.assertIn(("E3", "two_slips"), kinds(p))
		detail = next(c["detail"] for c in p["checks"] if c["kind"] == "two_slips")
		self.assertEqual(detail, "On salary slips S3, S5")

	def test_a_worker_paid_in_two_entries_is_not_a_check(self):
		# A range download spans several weekly entries; one slip a week is normal.
		p = pack(lines=[line("P1", "E1", "Alice", "Vale", "Picking", 1, 1, 100, 100.0),
				line("P2", "E1", "Alice", "Vale", "Picking", 1, 1, 100, 100.0)],
			slips=[slip("PE1", "S1", "E1", "Alice", "P1", 100.0), slip("PE2", "S2", "E1", "Alice", "P2", 100.0)],
			days={"E1": 2})
		self.assertEqual(p["checks"], [])
		self.assertEqual(by_emp(p["workers"])["E1"]["flag"], "")

	def test_agreeing_amounts_within_a_cent_are_not_flagged(self):
		p = pack(lines=[line("P1", "E1", "Alice", "Vale", "Picking", 1, 1, 100, 100.0)],
			slips=[slip("PE1", "S1", "E1", "Alice", "P1", 100.004)], days={"E1": 1})
		self.assertEqual(p["checks"], [])
		self.assertEqual(by_emp(p["workers"])["E1"]["flag"], "")


class TestFarmPacks(unittest.TestCase):

	def test_a_two_farm_worker_appears_in_each_farm_pack_with_only_that_farms_lines(self):
		p = pack()
		v, e = filter_farm(p, "Vale"), filter_farm(p, "Endebess")
		self.assertIn("E1", by_emp(v["workers"]))
		self.assertIn("E1", by_emp(e["workers"]))
		self.assertTrue(all(r["farm"] == "Vale" for r in v["detail"]))
		self.assertEqual(by_emp(v["workers"])["E1"]["farms"], ["Vale"])

	def test_farm_pack_total_is_that_farms_line_total(self):
		p = pack()
		self.assertEqual(filter_farm(p, "Vale")["totals"]["slip_total"], 1150.0)
		self.assertEqual(filter_farm(p, "Endebess")["totals"]["slip_total"], 750.0)

	def test_farm_pack_checks_only_name_its_own_workers(self):
		v = filter_farm(pack(), "Vale")
		self.assertEqual({c["employee"] for c in v["checks"]}, {"E3"})
		self.assertNotIn("E9", by_emp(v["workers"]))

	def test_a_farm_pack_never_shows_another_farms_amounts(self):
		v = filter_farm(pack(days_farm={("E1", "Vale"): 2, ("E1", "Endebess"): 2}), "Vale")
		w = by_emp(v["workers"])["E1"]
		self.assertEqual((w["wm_amount"], w["slip_amount"], w["difference"], w["flag"], w["tasks"]),
			(200.0, None, None, None, 1))
		self.assertNotIn("E1", {c["employee"] for c in v["checks"]})
		e = {r["name"]: r for r in v["entries"]}
		self.assertEqual((e["PE1"]["workers"], e["PE1"]["amount"]), (2, 600.0))
		self.assertEqual(v["totals"]["by_farm"], {"Vale": {"workers": 2, "days": 7, "amount": 1150.0}})

		def walk(x):
			if isinstance(x, dict):
				for k, val in x.items():
					yield from walk(k)
					yield from walk(val)
			elif isinstance(x, (list, tuple)):
				for val in x:
					yield from walk(val)
			else:
				yield x
		seen = list(walk(v))
		self.assertNotIn("Endebess", seen)
		for secret in (500, 500.0, 300.0, 450.0, 750.0, 1400.0, 1900.0, 1950.0):
			self.assertNotIn(secret, seen)

	def test_farm_days_come_from_days_by_employee_farm(self):
		w = by_emp(filter_farm(pack(days_farm={("E1", "Vale"): 2}), "Vale")["workers"])
		self.assertEqual(w["E1"]["days_worked"], 2)
		self.assertEqual(w["E3"]["days_worked"], 5)  # all on Vale: falls back to whole-person days

	def test_a_two_farm_workers_farm_days_are_none_when_unknown(self):
		w = by_emp(filter_farm(pack(), "Vale")["workers"])
		self.assertIsNone(w["E1"]["days_worked"])
		self.assertEqual(filter_farm(pack(), "Vale")["totals"]["man_days"], 5)

	def test_farms_are_listed_sorted(self):
		self.assertEqual(farms_in(pack()), ["Endebess", "Vale"])


class TestByTask(unittest.TestCase):

	def test_by_task_groups_per_farm_and_task(self):
		rows = {(r["farm"], r["task"]): r for r in pack()["by_task"]}
		self.assertEqual(set(rows), {("Endebess", "Picking"), ("Vale", "Weeding"), ("Vale", "Pruning")})
		picking = rows[("Endebess", "Picking")]
		self.assertEqual((picking["workers"], picking["days"], picking["qty"], picking["amount"]),
			(2, 7, 75, 750.0))
		self.assertEqual(rows[("Vale", "Weeding")]["workers"], 2)


def sheet_rows(wb, name):
	return [list(r) for r in wb[name].iter_rows(values_only=True)]


def load(p):
	return openpyxl.load_workbook(io.BytesIO(workbook_bytes(p)))


class TestWorkbook(unittest.TestCase):

	def test_workbook_has_the_five_sheets_in_order(self):
		self.assertEqual(load(pack()).sheetnames, ["Payrolls", "Workers", "Detail", "By task", "Checks"])

	def test_workbook_workers_total_matches_the_pack(self):
		p = pack()
		rows = sheet_rows(load(p), "Workers")
		col = rows[0].index("Slip amount")
		self.assertEqual(sum(r[col] for r in rows[1:]), p["totals"]["slip_total"])

	def test_empty_checks_sheet_says_no_problems(self):
		p = pack(lines=[l for l in LINES if l["employee"] in ("E1",)],
			slips=[s for s in SLIPS if s["employee"] == "E1"])
		self.assertEqual(p["checks"], [])
		self.assertEqual(load(p)["Checks"]["A2"].value, "No problems found.")

	def test_checks_use_human_labels(self):
		rows = sheet_rows(load(pack()), "Checks")
		labels = {r[2] for r in rows[1:]}
		self.assertEqual(labels, {"Paid with no work behind it", "Slip and work amounts differ",
			"On more than one slip"})

	def test_headers_are_bold_frozen_and_filtered_on_every_sheet(self):
		wb = load(pack())
		for ws in wb:
			self.assertTrue(ws["A1"].font.bold, ws.title)
			self.assertEqual(ws.freeze_panes, "A2", ws.title)
			self.assertIsNotNone(ws.auto_filter.ref, ws.title)

	def test_money_columns_are_formatted_and_dates_are_real_dates(self):
		wb = load(pack())
		rows = sheet_rows(wb, "Detail")
		head = rows[0]
		amount = wb["Detail"].cell(2, head.index("Amount") + 1)
		self.assertEqual(amount.number_format, "#,##0.00")
		frm = wb["Detail"].cell(2, head.index("Work from") + 1)
		self.assertEqual(frm.value, datetime.datetime(2026, 10, 12))
		self.assertEqual(frm.number_format, "yyyy-mm-dd")
		self.assertEqual(wb["Payrolls"]["B2"].value, datetime.datetime(2026, 10, 12))

	def test_farms_are_comma_joined_in_the_workers_sheet(self):
		rows = sheet_rows(load(pack()), "Workers")
		self.assertEqual(rows[1][2], "Endebess, Vale")

	def test_farm_pack_renders_none_as_empty_cells(self):
		p = filter_farm(pack(days_farm={("E1", "Vale"): 2}), "Vale")
		rows = sheet_rows(load(p), "Workers")
		alice = next(r for r in rows if r[0] == "E1")
		self.assertEqual(alice[rows[0].index("Slip amount")], None)
		self.assertEqual(alice[rows[0].index("Flag")], None)
		self.assertEqual(alice[rows[0].index("WM amount")], 200)
		for r in rows:
			self.assertNotIn("None", [c for c in r if isinstance(c, str)])

	def test_flag_column_shows_human_labels(self):
		rows = sheet_rows(load(pack()), "Workers")
		flag = {r[0]: r[rows[0].index("Flag")] for r in rows[1:]}
		self.assertEqual(flag["E2"], "Slip and work amounts differ")
		self.assertEqual(flag["E3"], "On more than one slip")
		self.assertEqual(flag["E9"], "Paid with no work behind it")
		self.assertIn(flag["E1"], (None, ""))
		# Two checks on one worker are comma-joined labels. E3 paid 600 for 550 here.
		slips = [dict(s, wm_amount=600) if s["slip"] == "S4" else s for s in SLIPS]
		rows = sheet_rows(load(pack(slips=slips)), "Workers")
		flag = {r[0]: r[rows[0].index("Flag")] for r in rows[1:]}
		self.assertEqual(flag["E3"], "Slip and work amounts differ, On more than one slip")

	def test_text_that_looks_like_a_formula_is_stored_as_text(self):
		# A name or task typed as "=HYPERLINK(...)" must not run in an approver's Excel.
		evil = ["=HYPERLINK(\"http://x.test\",\"Click\")", "+1+1", "-2+3", "@SUM(A1)"]
		lines = [dict(LINES[0], task=t, block=t) for t in evil]
		p = pack(lines=lines, slips=[SLIPS[0]])
		ws = load(p)["Detail"]
		head = [c.value for c in ws[1]]
		cells = [ws.cell(i + 2, head.index("Task") + 1) for i in range(len(evil))]
		self.assertEqual([c.value for c in cells], evil)
		for c in cells:
			self.assertEqual(c.data_type, "s", c.value)
		# Numbers stay numbers, negative ones included.
		self.assertEqual(ws.cell(2, head.index("Amount") + 1).data_type, "n")

	def test_control_characters_are_dropped_not_fatal(self):
		lines = [dict(LINES[0], task="Pick\x0bing\x00", employee_name="Al\x1bice")]
		p = pack(lines=lines, slips=[dict(SLIPS[0], employee_name="Al\x1bice")])
		wb = load(p)
		rows = sheet_rows(wb, "Detail")
		self.assertEqual(rows[1][rows[0].index("Task")], "Picking")
		self.assertEqual(rows[1][rows[0].index("Name")], "Alice")

	def test_internal_keys_are_never_written(self):
		wb = load(pack(days_farm={("E1", "Vale"): 2}))
		self.assertEqual(len(wb.sheetnames), 5)
		text = " ".join(str(c) for ws in wb for row in ws.iter_rows(values_only=True) for c in row)
		self.assertNotIn("days_by_employee", text)
		self.assertNotIn("wmpay_entry", text)


class TestSubjectAndName(unittest.TestCase):

	def test_subject_formats(self):
		self.assertEqual(subject("Kaitet Task pay 28th sept", 711428.9),
			"Task pay for approval \u2014 Kaitet Task pay 28th sept \u2014 KES 711,429")
		self.assertTrue(subject("X", 1, revised=True, test=True).startswith("TEST: Revised: "))
		self.assertIn("\u2014 Vale \u2014 KES", subject("X", 1, farm="Vale"))

	def test_attachment_name(self):
		self.assertEqual(attachment_name("2026-09-28", "2026-10-04", "Vale"),
			"Task-pay-2026-09-28-to-2026-10-04-Vale.xlsx")
		self.assertEqual(attachment_name(datetime.date(2026, 9, 28), datetime.date(2026, 10, 4)),
			"Task-pay-2026-09-28-to-2026-10-04.xlsx")


class TestEmailBody(unittest.TestCase):

	def body(self, p=None, **kw):
		return email_html(p or pack(), "PE1", "https://x.test/app/payroll-entry/PE1", **kw)

	def test_names_are_escaped_in_the_body(self):
		lines = [dict(LINES[0], task="Pruning & <weeding>", farm="Vale & <Co>")]
		p = pack(lines=lines, slips=[SLIPS[0]])
		html = self.body(p, farms_without_fm=["Odd <farm>"])
		self.assertNotIn("<weeding>", html)
		self.assertNotIn("<Co>", html)
		self.assertNotIn("<farm>", html)
		self.assertIn("Vale &amp; &lt;Co&gt;", html)

	def test_body_for_a_full_payroll_stays_under_100kb(self):
		lines, slips = [], []
		for i in range(400):
			for j in range(4):
				lines.append(line(f"P{i}-{j}", f"E{i}", f"Worker {i}", f"Farm{j}", "Picking", 1, 1, 10, 10))
			slips.append(slip("PE1", f"S{i}", f"E{i}", f"Worker {i}", f"P{i}-0", 40))
		html = self.body(pack(lines=lines, slips=slips, days={}))
		self.assertLess(len(html.encode()), 100 * 1024)

	def test_body_shows_line_total_only_when_it_differs(self):
		self.assertIn("1,900", self.body())
		agree = pack(lines=[LINES[0]], slips=[SLIPS[0]])
		html = self.body(agree)
		self.assertNotIn("Work lines add up to", html)
		self.assertIn("Work lines add up to", self.body())

	def test_body_counts_problems_or_says_none(self):
		self.assertIn("3 problems found", self.body())
		self.assertIn("see the Checks sheet", self.body())
		agree = pack(lines=[LINES[0]], slips=[SLIPS[0]])
		self.assertIn("No problems found", self.body(agree))

	def test_farm_body_says_its_checks_cover_only_sole_farm_workers(self):
		p = filter_farm(pack(), "Vale")
		self.assertEqual(len(p["checks"]), 1)
		html = self.body(p, farm="Vale")
		self.assertIn("1 problem found among workers only on this farm</b> \u2014 see the Checks sheet.", html)
		two = filter_farm(pack(lines=LINES + [line("P6", "E9", "Zed", "Vale", "Weeding", 1, 1, 10, 10)]), "Vale")
		self.assertEqual(len(two["checks"]), 2)
		self.assertIn("2 problems found among workers only on this farm</b>", self.body(two, farm="Vale"))
		clean = filter_farm(pack(lines=[LINES[0]], slips=[SLIPS[0]]), "Endebess")
		self.assertIn("No problems found among workers only on this farm.", self.body(clean, farm="Endebess"))
		# The full pack keeps its wording.
		self.assertNotIn("only on this farm", self.body())

	def test_body_has_link_period_and_per_farm_table(self):
		html = self.body()
		self.assertIn('href="https://x.test/app/payroll-entry/PE1"', html)
		self.assertIn("Open the Payroll Entry", html)
		self.assertIn("2026-10-12", html)
		self.assertIn("2026-10-25", html)
		self.assertIn("Endebess", html)

	def test_farm_body_has_no_farm_table_or_other_farms(self):
		html = self.body(filter_farm(pack(), "Vale"), farm="Vale")
		self.assertNotIn("Endebess", html)

	def test_farms_without_a_manager_are_listed(self):
		self.assertIn("Saboti", self.body(farms_without_fm=["Saboti"]))
		self.assertNotIn("no farm manager", self.body().lower())


class TestRecipients(unittest.TestCase):

	def test_farm_recipients_skip_scopeless_and_disabled_and_other_stages(self):
		rows = [{"stage_label": "Actuals: Farm Manager", "scope": "Vale", "user": "kibett@x"},
			{"stage_label": "Actuals: Farm Manager", "scope": None, "user": "koskey@x"},
			{"stage_label": "Planner: Farm Approval", "scope": "Vale", "user": "tarus@x"},
			{"stage_label": "Actuals: Farm Manager", "scope": "Saboti", "user": "gone@x"}]
		self.assertEqual(farm_recipients(rows, {"kibett@x", "koskey@x", "tarus@x"}), {"Vale": ["kibett@x"]})

	def test_farm_recipients_do_not_repeat_a_user(self):
		rows = [{"stage_label": "Actuals: Farm Manager", "scope": "Vale", "user": "k@x"}] * 2
		self.assertEqual(farm_recipients(rows, {"k@x"}), {"Vale": ["k@x"]})

	def test_split_addresses(self):
		self.assertEqual(split_addresses("a@x\n b@x, a@x\n\n"), ["a@x", "b@x"])
		self.assertEqual(split_addresses(None), [])


class TestTaskNamesAndIds(unittest.TestCase):
	"""Lines carry the Task's name in `task` and its ID in `task_id`. Two tasks
	can share a name, so grouping goes by ID and the name is what is shown."""

	LINES_WITH_IDS = [
		dict(line("P1", "E1", "Alice", "Vale", "Weeding", 2, 20, 10, 200), task_id="TASK-1"),
		dict(line("P1", "E2", "Bob", "Vale", "Weeding", 3, 30, 10, 300), task_id="TASK-2"),
		dict(line("P1", "E2", "Bob", "Vale", "Weeding", 1, 10, 10, 100), task_id="TASK-2"),
	]
	SLIPS_WITH_IDS = [slip("PE1", "S1", "E1", "Alice", "P1", 200), slip("PE1", "S2", "E2", "Bob", "P1", 400)]

	def p(self):
		return pack(lines=self.LINES_WITH_IDS, slips=self.SLIPS_WITH_IDS)

	def test_by_task_groups_by_task_id_and_shows_the_name(self):
		rows = self.p()["by_task"]
		self.assertEqual([(r["task_id"], r["task"], r["amount"]) for r in rows],
			[("TASK-1", "Weeding", 200.0), ("TASK-2", "Weeding", 400.0)])

	def test_worker_task_count_goes_by_id(self):
		lines = self.LINES_WITH_IDS[:2] + [dict(self.LINES_WITH_IDS[2], task_id="TASK-3")]
		p = pack(lines=lines, slips=self.SLIPS_WITH_IDS)
		self.assertEqual(by_emp(p["workers"])["E2"]["tasks"], 2)
		self.assertEqual(by_emp(filter_farm(p, "Vale")["workers"])["E2"]["tasks"], 2)

	def test_by_task_falls_back_to_the_task_when_there_is_no_id(self):
		rows = {(r["farm"], r["task"]): r for r in pack()["by_task"]}
		self.assertEqual(rows[("Vale", "Weeding")]["task_id"], "Weeding")

	def test_detail_sheet_has_task_id_right_after_task(self):
		rows = sheet_rows(load(self.p()), "Detail")
		head = rows[0]
		self.assertEqual(head.index("Task ID"), head.index("Task") + 1)
		self.assertEqual(rows[1][head.index("Task ID")], "TASK-1")
		self.assertEqual(rows[1][head.index("Task")], "Weeding")

	def test_farm_pack_by_task_also_goes_by_id(self):
		self.assertEqual(len(filter_farm(self.p(), "Vale")["by_task"]), 2)
		self.assertEqual(by_emp(filter_farm(pack(), "Vale")["workers"])["E1"]["tasks"], 1)


if __name__ == "__main__":
	unittest.main()
